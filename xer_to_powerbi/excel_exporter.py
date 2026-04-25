"""
Writes the 5 output sheets to an Excel workbook with basic formatting.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from pathlib import Path


# ── Palette ───────────────────────────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F3864")   # dark navy
ALT_ROW_FILL = PatternFill("solid", fgColor="EBF0FA")   # light blue
WHITE_FILL   = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
BODY_FONT    = Font(name="Calibri", size=10)
CENTER       = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT         = Alignment(horizontal="left",   vertical="center")

THIN_BORDER  = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)

# Status colour coding (for Activities sheet)
STATUS_COLORS = {
    "Completed":   "70AD47",
    "In Progress": "FFC000",
    "Not Started": "BDD7EE",
}

# Number formats
FMT_PCT      = "0.00%"
FMT_PCT_INT  = '0.00"%"'
FMT_SAR      = '#,##0.00 "SAR"'
FMT_INT      = "#,##0"
FMT_DATE     = "YYYY-MM-DD"
FMT_DECIMAL  = "#,##0.00"


def _apply_header(ws, row_num: int = 1) -> None:
    for cell in ws[row_num]:
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER
        cell.border    = THIN_BORDER


def _apply_body(ws, start_row: int = 2, status_col_idx: int | None = None) -> None:
    for r_idx, row in enumerate(ws.iter_rows(min_row=start_row), start=start_row):
        alt = (r_idx % 2 == 0)
        fill = ALT_ROW_FILL if alt else WHITE_FILL
        for cell in row:
            cell.font      = BODY_FONT
            cell.fill      = fill
            cell.alignment = LEFT
            cell.border    = THIN_BORDER

    # Status colour override
    if status_col_idx is not None:
        for row in ws.iter_rows(min_row=2, min_col=status_col_idx,
                                max_col=status_col_idx):
            for cell in row:
                colour = STATUS_COLORS.get(str(cell.value))
                if colour:
                    cell.fill = PatternFill("solid", fgColor=colour)
                    cell.alignment = CENTER


def _auto_width(ws, min_w: int = 10, max_w: int = 40) -> None:
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = \
            max(min_w, min(max_len + 2, max_w))


def _freeze(ws, cell: str = "A2") -> None:
    ws.freeze_panes = cell


def _write_df(ws, df: pd.DataFrame) -> None:
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)


# ── Column format maps ────────────────────────────────────────────────────────

_ACTIVITIES_FMTS = {
    "planned_start":       FMT_DATE,
    "planned_finish":      FMT_DATE,
    "actual_start":        FMT_DATE,
    "actual_finish":       FMT_DATE,
    "original_duration":   FMT_DECIMAL,
    "remaining_duration":  FMT_DECIMAL,
    "phys_complete_pct":   FMT_DECIMAL,
    "planned_pct":         FMT_DECIMAL,
    "variance_pct":        FMT_DECIMAL,
    "weight":              "#,##0.0000",
}

_RESOURCES_FMTS = {
    "unit_price":    FMT_SAR,
    "target_qty":    FMT_DECIMAL,
    "act_qty":       FMT_DECIMAL,
    "remain_qty":    FMT_DECIMAL,
    "target_cost":   FMT_SAR,
    "act_cost":      FMT_SAR,
    "ev_cost":       FMT_SAR,
    "cv":            FMT_SAR,
}

_WBS_FMTS = {
    "total_planned_cost": FMT_SAR,
    "total_actual_cost":  FMT_SAR,
    "wbs_level":          FMT_INT,
}

_PROJ_FMTS = {
    "data_date":       FMT_DATE,
    "planned_start":   FMT_DATE,
    "planned_finish":  FMT_DATE,
    "total_activities":       FMT_INT,
    "completed_activities":   FMT_INT,
    "inprogress_activities":  FMT_INT,
    "overall_planned_pct":    FMT_DECIMAL,
    "overall_actual_pct":     FMT_DECIMAL,
    "overall_variance_pct":   FMT_DECIMAL,
    "BAC":  FMT_SAR,
    "PV":   FMT_SAR,
    "EV":   FMT_SAR,
    "AC":   FMT_SAR,
    "SPI":  FMT_DECIMAL,
    "CPI":  FMT_DECIMAL,
    "EAC":  FMT_SAR,
    "VAC":  FMT_SAR,
}

_SCURVE_FMTS = {
    "period_date":        FMT_DATE,
    "planned_cum_pct":    FMT_DECIMAL,
    "actual_cum_pct":     FMT_DECIMAL,
    "planned_cum_cost":   FMT_SAR,
    "actual_cum_cost":    FMT_SAR,
}


def _apply_col_formats(ws, df: pd.DataFrame, fmts: dict[str, str]) -> None:
    col_indices = {name: idx + 1 for idx, name in enumerate(df.columns)}
    for col_name, fmt in fmts.items():
        col_idx = col_indices.get(col_name)
        if col_idx is None:
            continue
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter][1:]:   # skip header
            cell.number_format = fmt


# ── Public API ────────────────────────────────────────────────────────────────

def export_excel(datasets: dict[str, pd.DataFrame], output_path: str | Path) -> Path:
    """Write all 5 sheets to an Excel workbook."""
    output_path = Path(output_path)

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    _write_sheet_activities(wb, datasets.get("activities", pd.DataFrame()))
    _write_sheet_wbs(wb,        datasets.get("wbs",          pd.DataFrame()))
    _write_sheet_resources(wb,  datasets.get("resources",    pd.DataFrame()))
    _write_sheet_project(wb,    datasets.get("project_info", pd.DataFrame()))
    _write_sheet_scurve(wb,     datasets.get("scurve",       pd.DataFrame()))

    wb.save(output_path)
    print(f"[excel_exporter] Saved → {output_path}")
    return output_path


def _write_sheet_activities(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Activities")
    if df.empty:
        ws.append(["No data"])
        return

    _write_df(ws, df)
    _apply_header(ws)

    # Find status column
    status_idx = None
    if "status" in df.columns:
        status_idx = list(df.columns).index("status") + 1

    _apply_body(ws, status_col_idx=status_idx)
    _apply_col_formats(ws, df, _ACTIVITIES_FMTS)
    _auto_width(ws)
    _freeze(ws, "C2")
    ws.auto_filter.ref = ws.dimensions


def _write_sheet_wbs(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("WBS")
    if df.empty:
        ws.append(["No data"])
        return

    _write_df(ws, df)
    _apply_header(ws)
    _apply_body(ws)
    _apply_col_formats(ws, df, _WBS_FMTS)
    _auto_width(ws)
    _freeze(ws)
    ws.auto_filter.ref = ws.dimensions


def _write_sheet_resources(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Resources")
    if df.empty:
        ws.append(["No data"])
        return

    _write_df(ws, df)
    _apply_header(ws)
    _apply_body(ws)
    _apply_col_formats(ws, df, _RESOURCES_FMTS)
    _auto_width(ws)
    _freeze(ws, "C2")
    ws.auto_filter.ref = ws.dimensions


def _write_sheet_project(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Project_Info")
    if df.empty:
        ws.append(["No data"])
        return

    _write_df(ws, df)
    _apply_header(ws)
    _apply_body(ws)
    _apply_col_formats(ws, df, _PROJ_FMTS)
    _auto_width(ws)


def _write_sheet_scurve(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("SCurve")
    if df.empty:
        ws.append(["No data"])
        return

    _write_df(ws, df)
    _apply_header(ws)
    _apply_body(ws)
    _apply_col_formats(ws, df, _SCURVE_FMTS)
    _auto_width(ws)
    _freeze(ws)
    ws.auto_filter.ref = ws.dimensions
